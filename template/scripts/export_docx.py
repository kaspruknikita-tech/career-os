#!/usr/bin/env python3
"""Convert a tailored resume from Markdown to a clean DOCX.

Usage:
    python export_docx.py resume.md resume.docx

Expected Markdown shape:
    # Name
    contact line (email | phone | linkedin ...)

    ## Section Heading

    **Company** | Location | Dates
    - bullet
    - bullet

Requires: pip install python-docx
"""
import re
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_bottom_border(paragraph, size=4, color="BFBFBF"):
    """Add a thin horizontal rule under a paragraph (used for section dividers)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def remove_bottom_border(p_pr):
    """Drop a style's built-in bottom border (Word's Title style has one)."""
    for bdr in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(bdr)


def style_headings(doc):
    """Flatten heading colors to black and give section headings a divider line."""
    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(3)
    # Word's Title style ships with its own underline: it splits name from contacts,
    # so the header reads as two blocks. One rule under the whole header is enough.
    remove_bottom_border(title.paragraph_format.element)

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Arial"
    h1.font.size = Pt(13)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Arial"
    h2.font.size = Pt(11.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(2)


def add_title(doc, text):
    """Name in black, the title after the slash in grey regular: one line, two weights."""
    p = doc.add_heading("", level=0)
    name, sep, role = text.partition(" / ")
    run = p.add_run(name)
    run.bold = True
    if role:
        sep_run = p.add_run("  /  ")
        sep_run.bold = False
        sep_run.font.color.rgb = RGBColor(0xB0, 0xB0, 0xB0)
        role_run = p.add_run(role)
        role_run.bold = False
        role_run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    return p


def style_contact_line(paragraph):
    """Contact row: smaller, grey, tight under the name."""
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    paragraph.paragraph_format.space_after = Pt(0)


def add_bold_italic_runs(paragraph, text):
    """Split on **bold** and *italic* markers and add runs accordingly."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    for chunk in pattern.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            paragraph.add_run(chunk)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.05
    style_headings(doc)

    list_bullet = doc.styles["List Bullet"]
    list_bullet.font.name = "Arial"
    list_bullet.font.size = Pt(10.5)
    list_bullet.paragraph_format.space_after = Pt(3)
    list_bullet.paragraph_format.line_spacing = 1.05
    list_bullet.paragraph_format.left_indent = Inches(0.22)

    in_header = False       # between the name heading and the first real section
    header_done = False
    last_header_p = None    # last paragraph of the header block, gets the divider
    for raw in lines:
        line = raw.rstrip()

        if not line or line.strip() == "---":
            continue

        if line.startswith("# "):
            add_title(doc, line[2:].strip())
            in_header = True
            continue

        if line.startswith("## "):
            # "## Role Title" directly under the name is a subtitle, not a section:
            # no divider, so the header does not read as its own chapter.
            if in_header and last_header_p is None:
                p = doc.add_paragraph()
                run = p.add_run(line[3:].strip())
                run.bold = True
                run.font.size = Pt(13)
                p.paragraph_format.space_after = Pt(6)
                last_header_p = p
                continue

            # first real section closes the header: divider goes under the whole block
            if in_header and not header_done:
                if last_header_p is not None:
                    add_bottom_border(last_header_p)
                    last_header_p.paragraph_format.space_after = Pt(12)
                in_header = False
                header_done = True

            p = doc.add_heading(line[3:].strip(), level=1)
            add_bottom_border(p)
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            continue

        if line.strip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_bold_italic_runs(p, line.strip()[2:].strip())
            continue

        p = doc.add_paragraph()
        add_bold_italic_runs(p, line.strip())

        # header lines stay tight together, the divider is added once at the end
        if in_header:
            style_contact_line(p)
            last_header_p = p

    doc.save(docx_path)
    print(f"saved {docx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
